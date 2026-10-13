"""Cross-platform EXE smoke test: start binary, upload Unicode-named DOCX, run, check viewer+Excel."""
from __future__ import annotations

import hashlib
from io import BytesIO
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import time
import urllib.parse
import urllib.request
from zipfile import ZipFile

HEALTH_TIMEOUT = 60.0
RUN_TIMEOUT = 60.0


def _force_utf8_stdio() -> None:
    # Windows CI console defaults to a legacy code page (e.g. cp1252):
    # printing the Persian stored name would crash with UnicodeEncodeError.
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")
        except Exception:
            pass


_force_utf8_stdio()

# All smoke-test requests are local, including on machines with a system proxy.
_HTTP = urllib.request.build_opener(urllib.request.ProxyHandler({}))


def _read(base: str, path: str, timeout: float = 10.0):
    req = urllib.request.Request(f"{base}{path}")
    with _HTTP.open(req, timeout=timeout) as r:
        return r.status, dict(r.headers), r.read()


def _multipart(items: list[tuple[str, bytes]], field: str = "files") -> tuple[bytes, str]:
    # Like a real browser folder upload: the Persian webkitRelativePath
    # ("batch/رزومه فارسی.docx") as the filename, UTF-8 encoded.
    boundary = "----SmokeB1"
    parts = []
    for relpath, data in items:
        head = (
            f'--{boundary}\r\nContent-Disposition: form-data; name="{field}"; '
            f'filename="{relpath}"\r\n'
            "Content-Type: application/octet-stream\r\n\r\n"
        ).encode("utf-8")
        parts.append(head + data + b"\r\n")
    return b"".join(parts) + f"--{boundary}--\r\n".encode("ascii"), boundary


def _docx_bytes(candidate: str, large: bool = False) -> bytes:
    from docx import Document
    from io import BytesIO

    doc = Document()
    doc.add_paragraph(candidate)
    doc.add_paragraph("Over 7 years in software development.")
    doc.add_paragraph("React Python Docker")
    buf = BytesIO()
    doc.save(buf)
    if large:
        from zipfile import ZipFile, ZIP_STORED
        # A valid DOCX over 1 MB exercises multipart's disk-backed temporary file.
        with ZipFile(buf, "a") as archive:
            archive.writestr("word/media/spool-test.bin", b"x" * (2 * 1024 * 1024), compress_type=ZIP_STORED)
    return buf.getvalue()


def _pdf_bytes(candidate: str) -> bytes:
    from io import BytesIO
    from reportlab.pdfgen import canvas

    buf = BytesIO()
    pdf = canvas.Canvas(buf)
    for page in range(1, 4):
        pdf.drawString(60, 740, f"{candidate} - page {page}")
        pdf.drawString(60, 710, "Over 7 years of experience in software development. React Python Docker")
        pdf.showPage()
    pdf.save()
    return buf.getvalue()


def _check_pdf_assets(base: str) -> None:
    frontend = Path(__file__).resolve().parents[1] / "source" / "frontend"
    vendor = frontend / "vendor" / "pdfjs"
    assets = [frontend / "js" / "pdf_preview.mjs"] + [p for p in vendor.rglob("*") if p.is_file()]
    assert (vendor / "pdf.worker.min.mjs").is_file(), "PDF worker missing from build source"
    for path in assets:
        relative = path.relative_to(frontend).as_posix()
        status, headers, data = _read(base, "/" + urllib.parse.quote(relative))
        assert status == 200 and hashlib.sha256(data).digest() == hashlib.sha256(path.read_bytes()).digest(), relative
        media = next((value for key, value in headers.items() if key.lower() == "content-type"), "").split(";", 1)[0]
        if path.suffix == ".mjs":
            assert media == "text/javascript", (relative, media)
        elif path.suffix == ".wasm":
            assert media == "application/wasm", (relative, media)
    print(f"packaged PDF assets OK: {len(assets)} files, matching worker/modules, JavaScript/WASM MIME types", flush=True)


def _instance_port(env: dict) -> int | None:
    if os.name == "nt":
        base = env.get("LOCALAPPDATA") or str(os.path.expanduser("~\\AppData\\Local"))
        inst = os.path.join(base, "ResumeCheck", "instance.json")
    elif sys.platform == "darwin":
        inst = os.path.expanduser("~/Library/Application Support/ResumeCheck/instance.json")
    else:
        xdg = env.get("XDG_DATA_HOME") or os.path.expanduser("~/.local/share")
        inst = os.path.join(xdg, "resume-check", "instance.json")
    try:
        with open(inst, encoding="utf-8") as f:
            return json.load(f)["port"]
    except Exception:
        return None


def _resolve_exe(exe: str) -> str:
    # macOS CI passes the .app bundle; the real binary lives inside it.
    if exe.endswith(".app"):
        inner = os.path.join(exe, "Contents", "MacOS", "ResumeCheck")
        if os.path.isfile(inner):
            return inner
    return exe


def _check_bundled_modules(exe: str) -> None:
    """Confirm the frozen app really carries the Jobinja modules, not just their names.

    A missing hidden import only shows up when Jobinja is used, which is never
    exercised offline; CI would ship a broken Windows/Linux build unnoticed.
    """
    try:
        from PyInstaller.archive.readers import CArchiveReader
        from PyInstaller.loader.pyimod01_archive import ZlibArchiveReader
    except Exception:
        print("bundled module check skipped (PyInstaller not installed)", flush=True)
        return
    required = {"backend.app.jobinja", "backend.app.jobinja_challenge", "backend.app.jobinja_browser",
                "backend.app.report", "backend.app.runner", "backend.app.routes_runs"}
    archive = CArchiveReader(exe)
    with tempfile.TemporaryDirectory(prefix="resume-check-pyz-") as work:
        target = os.path.join(work, "PYZ.pyz")
        with open(target, "wb") as handle:
            handle.write(archive.extract("PYZ.pyz"))
        modules = set(ZlibArchiveReader(target).toc)
    missing = sorted(required - modules)
    assert not missing, f"frozen app is missing modules: {missing}"
    print(f"Packaged modules OK: {len(required)} backend modules bundled", flush=True)


def _check_zip_exports(base: str, run_id: str, state: dict, items: list) -> None:
    # Set all three statuses without depending on extraction heuristics.
    for result, decision in zip(state["results"], ("rejected", "needs-review")):
        req = urllib.request.Request(f"{base}/api/runs/{run_id}/review",
            data=json.dumps({"filename": result["filename"], "decision": decision}).encode(),
            headers={"Content-Type": "application/json"})
        with _HTTP.open(req, timeout=30) as response:
            state = json.load(response)
    statuses = {"all": None, "approved": "Qualified", "rejected": "Not Qualified", "needs-review": "Undetermined"}
    originals = {r["filename"]: (Path(state["folderPath"]) / r["filename"]).read_bytes() for r in state["results"]}
    for category, status_filter in statuses.items():
        expected = sum(status_filter is None or r["finalStatus"] == status_filter for r in state["results"])
        status, _, data = _read(base, f"/api/runs/{run_id}/export/{category}", timeout=30)
        assert status == 200
        with ZipFile(BytesIO(data)) as archive, tempfile.TemporaryDirectory(prefix="resume-check-unzip-") as extracted:
            names = archive.namelist()
            resumes = [name for name in names if "/resumes/" in name]
            assert len(resumes) == expected and len(names) == expected + 1, names
            assert len({name.casefold() for name in names}) == len(names), names
            assert archive.testzip() is None
            archive.extractall(extracted)
            assert sorted((Path(extracted) / name).read_bytes() for name in resumes) == sorted(
                originals[r["filename"]] for r in state["results"]
                if status_filter is None or r["finalStatus"] == status_filter)
            if category == "all":
                assert any("رزومه فارسی" in name for name in resumes), names
                assert sorted(hashlib.sha256((Path(extracted) / name).read_bytes()).hexdigest() for name in resumes) == sorted(
                    hashlib.sha256(content).hexdigest() for _, content in items)
            print(f"ZIP {category}: {expected} resumes + report, {len(data)} bytes, extracted and CRC validated; names={names}", flush=True)
    # Hide only synthetic copies temporarily; confirm the EXE refuses an incomplete ZIP.
    filename = state["results"][0]["filename"]
    source = Path(state["folderPath"]) / filename
    content = source.read_bytes()
    output = next(path for path in Path(state["outputPath"]).iterdir()
                  if path.is_file() and path.suffix == source.suffix and path.read_bytes() == content)
    moved = []
    try:
        for path in (source, output):
            backup = path.with_name(path.name + ".smoke-backup")
            path.rename(backup)
            moved.append((path, backup))
        try:
            _read(base, f"/api/runs/{run_id}/export/all")
            raise AssertionError("Missing resume was silently exported")
        except urllib.error.HTTPError as exc:
            detail = json.load(exc)["detail"]
            assert exc.code == 409, (exc.code, detail)
            # The refusal must name the count and the missing output, not a generic disk error.
            assert "1 فایل رزومه" in detail and "پوشهٔ خروجی" in detail, detail
            assert "فضای دیسک" not in detail, detail
    finally:
        for path, backup in moved:
            backup.rename(path)
    print("Packaged ZIP missing-file error OK", flush=True)


def _check_jobinja_import(base: str, output_folder: Path) -> None:
    """Exercise packaged Excel import and its failure report without external requests."""
    from openpyxl import Workbook, load_workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["#شناسه", "عنوان آگهی", "نام متقاضی", "شماره تماس متقاضی", "ایمیل",
                  "تاریخ ارسال درخواست همکاری", "آدرس فایل رزومه متقاضی"])
    sheet.append([1, "آگهی آزمایشی", "متقاضی آزمایشی", None, None, None, "مشاهده رزومه"])
    sheet["G2"].hyperlink = "https://example.invalid/not-jobinja"
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    body, boundary = _multipart([("جابینجا.xlsx", buffer.getvalue())], field="file")
    request = urllib.request.Request(f"{base}/api/jobinja/imports", data=body,
        headers={"Content-Type": f"multipart/form-data; boundary={boundary}"})
    with _HTTP.open(request, timeout=30) as response:
        imported = json.load(response)
    assert imported["total"] == 1 and imported["failed"] == 1 and imported["state"] == "ready"
    import_id = imported["importId"]
    request = urllib.request.Request(f"{base}/api/jobinja/imports/{import_id}/run",
        data=json.dumps({"criteria": {"minYears": 5}, "outputFolder": str(output_folder)}).encode(),
        headers={"Content-Type": "application/json"})
    with _HTTP.open(request, timeout=30) as response:
        json.load(response)
    deadline = time.time() + RUN_TIMEOUT
    while time.time() < deadline:
        _, _, raw = _read(base, f"/api/jobinja/imports/{import_id}")
        state = json.loads(raw)
        if state["runId"]:
            break
        assert state["state"] != "failed", state["error"]
        time.sleep(.1)
    assert state["runId"] and state["completed"] == 1 and state["downloaded"] == 0
    _, _, raw = _read(base, f"/api/runs/{state['runId']}")
    run = json.loads(raw)
    assert run["processed"] == 0 and not run["results"] and len(run["downloadFailures"]) == 1
    _, _, report = _read(base, f"/api/runs/{state['runId']}/report")
    workbook = load_workbook(BytesIO(report))
    assert workbook["دریافت رزومه‌ها"]["F2"].value == "دریافت نشد"
    assert workbook["دریافت رزومه‌ها"]["I2"].value is None
    workbook.close()
    status, _, font = _read(base, "/fonts/Vazirmatn.woff2")
    expected_font = Path(__file__).resolve().parents[1] / "source/frontend/fonts/Vazirmatn.woff2"
    assert status == 200 and font == expected_font.read_bytes() and font.startswith(b"wOF2")
    print("Packaged Jobinja XLSX/parser/failure report + bundled Persian font OK; no external downloads", flush=True)


def main(exe: str, zip_only: bool = False) -> int:
    exe = _resolve_exe(exe)
    test_data = tempfile.TemporaryDirectory(prefix="resume-check-smoke-")
    env = dict(os.environ, RESUMECHECK_HEADLESS="1", PYTHONUTF8="1", PYTHONIOENCODING="utf-8",
               LOCALAPPDATA=test_data.name, XDG_DATA_HOME=test_data.name)
    log = open("smoke-server.log", "ab")
    proc = subprocess.Popen(
        [exe], stdin=subprocess.DEVNULL, stdout=log, stderr=subprocess.STDOUT, env=env
    )
    try:
        deadline = time.time() + HEALTH_TIMEOUT
        port = None
        while time.time() < deadline:
            port = _instance_port(env)
            if port:
                try:
                    status, _, body = _read(f"http://127.0.0.1:{port}", "/api/health")
                    if status == 200 and json.loads(body).get("status") == "ok":
                        break
                except Exception:
                    pass
            if proc.poll() is not None:
                print("server exited early", flush=True)
                return 1
            time.sleep(0.5)
        else:
            print("health check timed out", flush=True)
            return 1

        base = f"http://127.0.0.1:{port}"
        _check_bundled_modules(exe)
        if not zip_only:
            status, _, html = _read(base, "/")
            assert status == 200 and "غربالگری رزومه".encode() in html
            _check_pdf_assets(base)
        fa_name = "رزومه فارسی.docx"
        names = [fa_name, f"تیم اول/{fa_name}", f"تیم دوم/{fa_name}",
                 "cv.docx", "CV.docx", "cv+.docx", "cv#.docx", "CON.docx",
                 "AUX/رزومه.docx", "first/a-CV.docx", "FIRST-a/cv.docx", "large.docx",
                 "رزومه فارسی چند صفحه.pdf", "Team Two/candidate with spaces.pdf"]
        if zip_only:
            names += [f"رزومه {i}.docx" for i in range(43 - len(names))]
        items = [(f"batch/{name}", _pdf_bytes(f"Smoke Candidate {i}") if name.endswith(".pdf")
                  else _docx_bytes(f"Smoke Candidate {i}", large=name == "large.docx"))
                 for i, name in enumerate(names)]
        body, boundary = _multipart(items)
        req = urllib.request.Request(
            f"{base}/api/uploads", data=body,
            headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
        )
        with _HTTP.open(req, timeout=30) as r:
            up = json.load(r)
        upload_id = up["uploadId"]
        assert up["fileCount"] == len(items), up

        output_folder = Path(test_data.name) / "نتایج آزمون"
        output_folder.mkdir()
        payload = json.dumps({"criteria": {"minYears": 5},
                              "outputFolder": None if zip_only and os.name == "nt" else str(output_folder)}).encode()
        req = urllib.request.Request(
            f"{base}/api/uploads/{upload_id}/run", data=payload,
            headers={"Content-Type": "application/json"},
        )
        with _HTTP.open(req, timeout=30) as r:
            run_id = json.load(r)["runId"]

        deadline = time.time() + RUN_TIMEOUT
        state = {}
        while time.time() < deadline:
            _, _, raw = _read(base, f"/api/runs/{run_id}")
            state = json.loads(raw)
            if state["state"] == "finished":
                break
            time.sleep(0.5)
        assert state.get("state") == "finished", state
        assert state["found"] == len(items) and state["processed"] == len(items), state
        assert state["qualified"] == len(items), state
        if zip_only:
            _check_zip_exports(base, run_id, state, items)
            _check_jobinja_import(base, output_folder)
            assert proc.poll() is None, "launcher exited during ZIP test"
            return 0
        viewed = []
        for result in state["results"]:
            fname = result["filename"]
            status, _, content = _read(base, f"/api/runs/{run_id}/resumes/{urllib.parse.quote(fname)}")
            assert status == 200, fname
            if fname.endswith(".pdf"):
                req = urllib.request.Request(f"{base}/api/runs/{run_id}/preview",
                    data=json.dumps({"filename": fname}).encode(), headers={"Content-Type": "application/json"})
                with _HTTP.open(req, timeout=30) as response:
                    assert response.headers.get("Content-Type") == "application/octet-stream"
                    assert response.headers.get("Content-Disposition") is None
                    assert response.read() == content, "Preview bytes differ from stored PDF"
            viewed.append(hashlib.sha256(content).hexdigest())
        assert sorted(viewed) == sorted(hashlib.sha256(data).hexdigest() for _, data in items)
        status, _, xlsx = _read(base, f"/api/runs/{run_id}/report")
        assert status == 200 and xlsx[:2] == b"PK", status
        from io import BytesIO
        from openpyxl import load_workbook
        rows = list(load_workbook(BytesIO(xlsx)).active.values)
        header = next(i for i, row in enumerate(rows) if row[0] == "نام فایل رزومه")
        assert len(rows[header + 1:]) == len(items)
        assert proc.poll() is None, "launcher exited during upload/run"
        print(f"smoke OK: {len(items)} resumes, >1 MB upload, viewer integrity, Persian Excel", flush=True)
        return 0
    except Exception as exc:
        import traceback
        print(f"SMOKE FAILED: {exc}", flush=True)
        traceback.print_exc()
        return 1
    finally:
        if os.name == "nt" and proc.poll() is None:
            # Windows onefile has a bootloader parent and a server child.
            subprocess.run(["taskkill", "/PID", str(proc.pid), "/T", "/F"],
                           stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=False)
        elif proc.poll() is None:
            proc.terminate()
        try:
            proc.wait(timeout=15)
        except Exception:
            proc.kill()
            proc.wait(timeout=5)
        # Preserve the windowless EXE's backend diagnostics before deleting test storage.
        server_log = Path(test_data.name) / ("ResumeCheck" if os.name == "nt" else "resume-check") / "logs" / "server.log"
        if server_log.is_file():
            log.write(server_log.read_bytes())
        log.close()
        test_data.cleanup()


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1], zip_only="--zip-only" in sys.argv[2:]))
