"""Synthetic XLSX/HTML/resumes only: no real applicant URLs are opened."""
from email.message import Message
from io import BytesIO
import logging
import gzip
import os
from types import SimpleNamespace
from pathlib import Path
import socket
import time
import urllib.error
from zipfile import ZipFile
import errno

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
import pytest

from backend.app import jobinja, jobinja_challenge, routes_runs, runner, uploads
from backend.app.main import app

URL = "https://jobinja.ir/resume/sample?token=SENSITIVE%2BToken&value=one+two&item=1&item=2"
CHALLENGE = b"""<html><script>
document.cookie = '__arcsjs=synthetic; Path=/; Secure';
document.cookie = '__arcsjsc=synthetic2; Path=/; Secure';
setTimeout(() => location.reload(), 10);
</script></html>"""


def excel(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "نام شیت دلخواه"
    sheet.append(["نام متقاضی", "اطلاعات دیگر", "رزومه"])
    for name, link, kind in rows:
        sheet.append([name, "", "مشاهده رزومه"])
        cell = sheet.cell(sheet.max_row, 3)
        if kind == "hyperlink":
            cell.hyperlink = link
        elif kind == "formula":
            cell.value = '=HYPERLINK("' + link.replace('"', '""') + '","مشاهده رزومه")'
        else:
            cell.value = link
    workbook.create_sheet("نادیده گرفته شود").append(["not an applicant"])
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def pdf(text):
    from reportlab.pdfgen import canvas
    buffer = BytesIO()
    page = canvas.Canvas(buffer)
    page.drawString(30, 700, text)
    page.save()
    return buffer.getvalue()


def docx(text):
    document = Document()
    document.add_paragraph("Synthetic Applicant")
    document.add_paragraph(text)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.mark.parametrize("kind", ["hyperlink", "formula", "text"])
def test_first_sheet_real_url_and_query_are_preserved(kind):
    applicants = jobinja.parse_excel(excel([("نام فارسی", URL, kind)]))
    assert len(applicants) == 1
    assert applicants[0].jobinja_url == URL
    assert applicants[0].name == "نام فارسی"
    assert applicants[0].excel_row == 2
    assert "SENSITIVE" not in repr(applicants[0])


def test_empty_rows_invalid_links_and_duplicate_names():
    rows = jobinja.parse_excel(excel([("مشترک", URL, "hyperlink"), (None, None, "text"),
                                     ("مشترک", URL, "formula"), ("نامعتبر", "javascript:alert(1)", "text")]))
    assert len(rows) == 3
    assert rows[0].applicant_id != rows[1].applicant_id
    assert rows[2].download_status == "failed"


def test_link_column_has_no_fixed_position():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["رزومه", "نام و نام خانوادگی", "وب‌سایت شخصی"])
    sheet.append(["مشاهده رزومه", "شخص آزمایشی", "https://example.org"])
    sheet["A2"].hyperlink = URL
    buffer = BytesIO()
    workbook.save(buffer)
    rows = jobinja.parse_excel(buffer.getvalue())
    assert rows[0].jobinja_url == URL and rows[0].name == "شخص آزمایشی"


def test_actual_export_headers_and_524_applicants():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["#شناسه", "عنوان آگهی", "نام متقاضی", "شماره تماس متقاضی", "ایمیل",
                  "تاریخ ارسال درخواست همکاری", "آدرس فایل رزومه متقاضی"])
    for index in range(524):
        sheet.append([index, "آگهی آزمایشی", f"متقاضی {index}", None, None, None, "مشاهده رزومه"])
        sheet.cell(index + 2, 7).hyperlink = f"https://panel.jobinja.ir/application/share/synthetic-{index}?signing=TEST%2BSign&valid_until=TEST"
    buffer = BytesIO()
    workbook.save(buffer)
    rows = jobinja.parse_excel(buffer.getvalue())
    assert len(rows) == 524
    assert all(a.name == f"متقاضی {i}" and a.excel_row == i + 2
               and a.jobinja_url == sheet.cell(i + 2, 7).hyperlink.target and a.download_status == "pending"
               for i, a in enumerate(rows))


@pytest.mark.parametrize("data", [b"not Excel", b"", b"x" * (jobinja.MAX_EXCEL_BYTES + 1)])
def test_invalid_excel_fails_safely(data):
    with pytest.raises(jobinja.JobinjaError):
        jobinja.parse_excel(data)


@pytest.mark.parametrize("url", ["http://jobinja.ir/cv", "https://jobinja.ir.evil.org/cv",
                                "https://127.0.0.1/cv", "https://u:p@jobinja.ir/cv",
                                "https://jobinja.ir:8443/cv", "javascript:alert(1)"])
def test_untrusted_urls_are_rejected(url):
    assert not jobinja.valid_url(url)


class Response(BytesIO):
    def __init__(self, data, media, headers=None):
        super().__init__(data)
        self.status = 200
        self.headers = Message()
        self.headers["Content-Type"] = media
        for key, value in (headers or {}).items():
            self.headers[key] = value

    def info(self):
        return self.headers


@pytest.fixture
def fake_http(monkeypatch):
    responses, requests = {}, []
    class Opener:
        def __init__(self, handlers):
            self.cookies = next(h for h in handlers if isinstance(h, jobinja.urllib.request.HTTPCookieProcessor))

        def open(self, request, timeout):
            self.cookies.http_request(request)
            requests.append(request)
            value = responses[request.full_url]
            value = value(request) if callable(value) else value
            if isinstance(value, Exception):
                if isinstance(value, urllib.error.HTTPError):
                    self.cookies.http_response(request, value)
                raise value
            response = Response(*value)
            self.cookies.http_response(request, response)
            return response
    monkeypatch.setattr(jobinja.urllib.request, "build_opener", lambda *args: Opener(args))
    monkeypatch.setattr(jobinja.socket, "getaddrinfo", lambda *args, **kwargs: [
        (socket.AF_INET, socket.SOCK_STREAM, 6, "", ("93.184.216.34", 443))])
    return responses, requests


@pytest.mark.parametrize("extension", [".pdf", ".docx"])
def test_html_download_without_login_preserves_both_urls(fake_http, extension):
    responses, requests = fake_http
    download_url = "https://jobinja.ir/download?key=SENSITIVE%2BToken&x=one+two"
    if extension == ".pdf":
        from reportlab.pdfgen import canvas
        buffer = BytesIO()
        pdf = canvas.Canvas(buffer)
        pdf.drawString(30, 700, "5+ years of professional experience")
        pdf.save()
        data = buffer.getvalue()
    else:
        data = docx("5+ years of professional experience")
    responses[URL] = ('<html><a href="/download?key=SENSITIVE%2BToken&amp;x=one+two"><span>دانلود رزومه</span></a></html>'.encode(), "text/html")
    responses[download_url] = (data, "application/octet-stream")
    assert jobinja.fetch_resume(URL) == (data, extension)
    assert [r.full_url for r in requests] == [URL, download_url]
    assert all(not r.has_header("Cookie") and not r.has_header("Authorization") for r in requests)
    assert "Mozilla/5.0" in requests[0].get_header("User-agent")
    assert "text/html" in requests[0].get_header("Accept")
    assert requests[0].get_header("Accept-language").startswith("fa-IR")
    assert requests[1].get_header("Referer") == URL


def test_actual_html_button_uses_jobinja_cdn_and_preserves_signed_query(fake_http):
    responses, requests = fake_http
    page = "https://panel.jobinja.ir/application/share/synthetic?signing=TEST%2BSign&valid_until=TEST"
    link = "https://mstorage2.jobinjacdn.com/records/files/uploads/documents/synthetic.docx?requester=TEST&X-Amz-Credential=TEST%2Fvalue&X-Amz-Signature=TEST"
    data = docx("5+ years of professional experience")
    responses[page] = (f'<html><a class="btn btn-default" href="{link.replace("&", "&amp;")}">دانلود</a></html>'.encode(), "text/html")
    responses[link] = (data, "application/octet-stream")
    assert jobinja.fetch_resume(page) == (data, ".docx")
    assert [r.full_url for r in requests] == [page, link]
    assert requests[1].get_header("Referer") == "https://panel.jobinja.ir/"
    assert not jobinja.valid_download_url("https://jobinjacdn.com.evil.org/cv.pdf")


def test_browser_session_cookies_survive_redirect_and_download_but_not_next_applicant(fake_http):
    responses, requests = fake_http
    page = "https://jobinja.ir/shared"
    link = "https://jobinja.ir/file"
    headers = Message()
    headers["Location"] = page
    headers["Set-Cookie"] = "session=synthetic; Path=/; Secure; HttpOnly"
    responses[URL] = urllib.error.HTTPError(URL, 302, "redirect", headers, BytesIO())
    responses[page] = (b'<a href="/file" download>Download</a>', "text/html")
    data = docx("5+ years of professional experience")
    responses[link] = (data, "application/octet-stream")
    for _ in range(2):
        assert jobinja.fetch_resume(URL) == (data, ".docx")
    assert [r.get_header("Cookie") for r in requests] == [
        None, "session=synthetic", "session=synthetic",
        None, "session=synthetic", "session=synthetic"]


def test_page_cookies_are_not_sent_to_cdn(fake_http):
    responses, requests = fake_http
    link = "https://mstorage2.jobinjacdn.com/file.docx"
    responses[URL] = (f'<a href="{link}" download>Download</a>'.encode(), "text/html",
                      {"Set-Cookie": "session=synthetic; Path=/; Secure"})
    data = docx("5+ years of professional experience")
    responses[link] = (data, "application/octet-stream")
    assert jobinja.fetch_resume(URL) == (data, ".docx")
    assert not requests[1].has_header("Cookie")
    assert requests[1].get_header("Referer") == "https://jobinja.ir/"


def test_gzip_html_and_resume_are_decoded(fake_http):
    responses, _ = fake_http
    link = "https://jobinja.ir/file"
    data = docx("5+ years of professional experience")
    responses[URL] = (gzip.compress(b'<a href="/file" download>Download</a>'), "text/html",
                      {"Content-Encoding": "gzip"})
    responses[link] = (gzip.compress(data), "application/octet-stream", {"Content-Encoding": "gzip"})
    assert jobinja.fetch_resume(URL) == (data, ".docx")


def test_gzip_expansion_obeys_html_size_limit(fake_http, monkeypatch):
    responses, _ = fake_http
    monkeypatch.setattr(jobinja, "MAX_HTML_BYTES", 128)
    responses[URL] = (gzip.compress(b" " * 129), "text/html", {"Content-Encoding": "gzip"})
    with pytest.raises(jobinja.JobinjaError, match="حجم"):
        jobinja.fetch_resume(URL)


def test_http_error_reports_status_without_signed_url(fake_http):
    responses, _ = fake_http
    responses[URL] = urllib.error.HTTPError(URL, 403, "SENSITIVE", Message(), BytesIO())
    with pytest.raises(jobinja.JobinjaError, match="HTTP 403") as error:
        jobinja.fetch_resume(URL)
    assert "SENSITIVE" not in str(error.value)


def test_missing_download_includes_safe_step_log(fake_http):
    responses, _ = fake_http
    responses[URL] = (b'<html><script>var token="SENSITIVE";</script><input type="password"></html>', "text/html")
    with pytest.raises(jobinja.JobinjaError) as error:
        jobinja.fetch_resume(URL)
    message = str(error.value)
    assert "گزارش مراحل دریافت" in message
    assert "HTTP 200" in message and "Content-Type=text/html" in message
    assert "لینک‌های دانلود: 0" in message
    assert "script=1" in message and "password=True" in message
    assert "SENSITIVE" not in message and "<html>" not in message and URL not in message


def test_network_failure_log_does_not_expose_exception_url(fake_http):
    responses, _ = fake_http
    responses[URL] = urllib.error.URLError(TimeoutError("SENSITIVE " + URL))
    with pytest.raises(jobinja.JobinjaError) as error:
        jobinja.fetch_resume(URL)
    assert "TimeoutError" in str(error.value)
    assert "SENSITIVE" not in str(error.value)


def test_waiting_page_is_recognized_and_browser_failure_is_actionable(fake_http, monkeypatch):
    # Playwright is optional; the browser fallback tests only run where it is installed.
    playwright_api = pytest.importorskip("playwright.sync_api", reason="Playwright is optional")
    BrowserType, Error = playwright_api.BrowserType, playwright_api.Error
    responses, _ = fake_http
    responses[URL] = (CHALLENGE, "text/html")
    assert jobinja._is_browser_challenge(CHALLENGE.decode())
    assert not jobinja._is_browser_challenge('<script src="resume.js"></script>')
    attempts = []
    def unavailable(browser_type, **kwargs):
        attempts.append(kwargs["channel"])
        raise Error("SENSITIVE browser path")
    monkeypatch.setattr(BrowserType, "launch", unavailable)
    with pytest.raises(jobinja.JobinjaError) as error:
        jobinja.fetch_resume(URL)
    assert attempts == ["chrome", "msedge"]
    assert "Chrome یا Microsoft Edge" in str(error.value)
    assert "__arcsjs" in str(error.value)
    assert "SENSITIVE" not in str(error.value)


def challenge_page(key=6, nested=True):
    """A synthetic waiting page with the same shape as Jobinja's obfuscated one."""
    first = r"(![]+[])[+!+[]]+(![]+[])[+[]]"
    second = r"(![]+[])[+!+[]]+(!![]+[])[!+[]+!+[]]"
    call = "E(E(" if nested else "E("
    return ("<html><script>var _0x1 = function (v, k) { if (!k) { k = 0x%x; } var out = '';"
            " for (var i = 0; i < v.length; i++) { out += String.fromCharCode(k ^ v.charCodeAt(i)); }"
            " return out; }; var E = _0x1;"
            " function values() { return {'value_v1': eval(\"%s\"), 'value': eval(\"%s\")}; }"
            " var hash_v1 = %s(values().value_v1); var hash = %s(values().value);"
            " document.addEventListener('DOMContentLoaded', function () {"
            " document.cookie = '__arcsjs=' + encodeURIComponent(hash_v1) + '; Max-Age=900; Path=/';"
            " document.cookie = '__arcsjsc=' + encodeURIComponent(hash) + '; Max-Age=900; Path=/';"
            " location.reload(); });</script></html>") % (key, first, second, call, call)


def test_challenge_solver_reads_obfuscated_values():
    assert jobinja_challenge.solve(challenge_page()) == {"__arcsjs": "af", "__arcsjsc": "au"}


def test_challenge_solver_honours_key_and_nesting():
    assert jobinja_challenge.solve(challenge_page(key=5, nested=False)) == {"__arcsjs": "dc", "__arcsjsc": "dp"}


def test_challenge_solver_rejects_an_unknown_page():
    with pytest.raises(jobinja_challenge.ChallengeError):
        jobinja_challenge.solve("<html><script>location.reload();</script></html>")


def test_waiting_page_is_solved_without_a_browser(fake_http, monkeypatch):
    from backend.app import jobinja_browser
    responses, requests = fake_http
    link = "https://jobinja.ir/user/NO-SYNTHETIC/download"
    data = docx("5+ years of professional experience")
    waiting = challenge_page().encode()
    page = '<html><a href="/user/NO-SYNTHETIC/download" download>دانلود</a></html>'.encode()
    responses[URL] = lambda request: (waiting if not request.has_header("Cookie") else page, "text/html")
    responses[link] = (data, "application/octet-stream")

    def unexpected(*args, **kwargs):
        raise AssertionError("the browser fallback must not be needed")

    monkeypatch.setattr(jobinja_browser, "render_resume_page", unexpected)
    assert jobinja.fetch_resume(URL) == (data, ".docx")
    assert [r.full_url for r in requests] == [URL, URL, link]
    assert "__arcsjs=af" in requests[1].get_header("Cookie")
    assert "__arcsjsc=au" in requests[2].get_header("Cookie")


def test_challenge_values_stay_out_of_the_trace(fake_http, monkeypatch):
    from backend.app import jobinja_browser
    responses, _ = fake_http
    responses[URL] = (challenge_page(key=5, nested=False).encode(), "text/html")
    monkeypatch.setattr(jobinja_browser, "render_resume_page", lambda *a, **k: ("<html></html>", URL))
    with pytest.raises(jobinja.JobinjaError) as error:
        jobinja.fetch_resume(URL)
    message = str(error.value)
    assert "__arcsjs" in message
    assert "dc" not in message and "dp" not in message


def test_waiting_page_again_falls_back_to_the_browser(fake_http, monkeypatch):
    from backend.app import jobinja_browser
    responses, requests = fake_http
    link = "https://jobinja.ir/file"
    data = docx("5+ years of professional experience")
    responses[URL] = (challenge_page().encode(), "text/html")
    responses[link] = (data, "application/octet-stream")
    used = []

    def render(url, cookies, trace):
        used.append(url)
        trace.append("مرورگر آزمایشی")
        return f'<a href="{link}" download>Download</a>', url

    monkeypatch.setattr(jobinja_browser, "render_resume_page", render)
    assert jobinja.fetch_resume(URL) == (data, ".docx")
    assert used == [URL]
    assert [r.full_url for r in requests] == [URL, URL, link]


@pytest.mark.skipif(os.environ.get("RESUMECHECK_BROWSER_TEST") != "1", reason="Requires installed Chrome/Edge")
@pytest.mark.parametrize("redirect", [None, "https://panel.jobinja.ir/after-waiting", "https://evil.example/blocked", "https://127.0.0.1/blocked"])
def test_real_browser_executes_waiting_script_and_checks_redirects(fake_http, monkeypatch, redirect):
    from playwright.sync_api import Route
    responses, requests = fake_http
    link = "https://jobinja.ir/download?token=SENSITIVE%2BToken"
    data = docx("5+ years of professional experience")
    responses[URL] = (CHALLENGE, "text/html")
    responses[link] = (data, "application/octet-stream")
    browser_requests = []
    real_fulfill = Route.fulfill

    def fetch(route, **kwargs):
        assert kwargs["max_redirects"] == 0
        target = kwargs["url"]
        browser_requests.append(target)
        headers = {"content-type": "text/html"}
        status = 200
        browser_cookies = {c["name"]: c["value"] for c in route.request.frame.page.context.cookies([target])}
        if browser_cookies.get("__arcsjs") != "synthetic":
            body = CHALLENGE
        elif redirect and target == URL:
            status, headers["location"], body = 302, redirect, b""
        else:
            body = f'<html><a href="{link.replace("&", "&amp;")}" download>Download</a></html>'.encode()
        return SimpleNamespace(status=status, headers=headers, body=lambda: body)

    def fulfill(route, *, response=None, **kwargs):
        if response is not None:
            kwargs = {"status": response.status, "headers": response.headers, "body": response.body()}
        return real_fulfill(route, **kwargs)

    monkeypatch.setattr(Route, "fetch", fetch)
    monkeypatch.setattr(Route, "fulfill", fulfill)
    if redirect and not jobinja.valid_download_url(redirect):
        with pytest.raises(jobinja.JobinjaError, match="آدرس دریافت رزومه معتبر نیست") as error:
            jobinja.fetch_resume(URL)
        assert redirect not in browser_requests
        assert "SENSITIVE" not in str(error.value)
    else:
        assert jobinja.fetch_resume(URL) == (data, ".docx")
        assert requests[-1].get_header("Cookie") == "__arcsjs=synthetic; __arcsjsc=synthetic2"
        assert requests[-1].get_header("Referer") == (URL if redirect is None else "https://panel.jobinja.ir/")
        if redirect:
            assert redirect in browser_requests
    assert browser_requests[:2] == [URL, URL]


def test_redirect_to_external_or_private_host_is_not_opened(fake_http):
    responses, requests = fake_http
    headers = Message()
    headers["Location"] = "https://127.0.0.1/private?token=SENSITIVE"
    responses[URL] = urllib.error.HTTPError(URL, 302, "redirect", headers, BytesIO())
    with pytest.raises(jobinja.JobinjaError):
        jobinja.fetch_resume(URL)
    assert len(requests) == 1


def test_invalid_file_is_not_staged(fake_http):
    responses, _ = fake_http
    responses[URL] = (b"pretend.pdf", "application/pdf")
    with pytest.raises(jobinja.JobinjaError, match="PDF یا DOCX معتبر"):
        jobinja.fetch_resume(URL)


def test_valid_encrypted_pdf_is_received_for_existing_manual_review(fake_http):
    from pypdf import PdfWriter
    responses, _ = fake_http
    pdf = PdfWriter()
    pdf.add_blank_page(width=200, height=200)
    pdf.encrypt("synthetic-password")
    buffer = BytesIO()
    pdf.write(buffer)
    data = buffer.getvalue()
    responses[URL] = (data, "application/pdf")
    assert jobinja.fetch_resume(URL) == (data, ".pdf")


@pytest.fixture
def isolated(tmp_path, monkeypatch):
    monkeypatch.setattr(routes_runs.tempfile, "gettempdir", lambda: str(tmp_path))
    monkeypatch.setattr(runner, "user_data_dir", lambda: tmp_path / "user-data")
    monkeypatch.setattr(jobinja, "_imports", {})
    for registry in ("_runs", "_uploads", "_output_files", "_output_targets"):
        monkeypatch.setattr(runner, registry, {})
    return TestClient(app)


def wait_for_import(client, import_id):
    for _ in range(100):
        state = client.get(f"/api/jobinja/imports/{import_id}").json()
        if state["runId"] or state["state"] == "failed":
            return state
        time.sleep(.02)
    raise AssertionError("Import did not finish")


def test_shared_pipeline_report_zip_and_no_token_logs(isolated, fake_http, caplog):
    client = isolated
    responses, requests = fake_http
    originals = []
    rows = []
    for index, text in enumerate(["2+ years of professional experience", "5+ years of professional experience",
                                  "7+ years of professional experience", "Jan 2010 - Dec 2020", "failure"]):
        url = URL + f"&candidate={index}"
        rows.append(("نام مشترک", url, "hyperlink"))
        if index < 4:
            download_url = "https://jobinja.ir/file?token=SENSITIVE&candidate=" + str(index)
            data = docx(text)
            originals.append(data)
            responses[url] = (f'<a href="{download_url.replace("&", "&amp;")}" download>دانلود</a>'.encode(), "text/html")
            responses[download_url] = (data, "application/octet-stream")
        else:
            responses[url] = urllib.error.HTTPError(url, 403, "SENSITIVE token denied", Message(), BytesIO())
    with caplog.at_level(logging.INFO):
        up = client.post("/api/jobinja/imports", files={"file": ("جابینجا.xlsx", excel(rows))})
        assert up.status_code == 200, up.text
        import_id = up.json()["importId"]
        assert not requests, "Uploading Excel must not start downloads"
        start = client.post(f"/api/jobinja/imports/{import_id}/run", json={"criteria": {"minYears": 5}})
        assert start.status_code == 200, start.text
        state = wait_for_import(client, import_id)
        assert state["downloaded"] == 4 and state["completed"] == 5 and state["failed"] == 1
        run_id = state["runId"]
        run = client.get(f"/api/runs/{run_id}").json()
        assert run["state"] == "finished" and run["processed"] == 4
        assert run["qualified"] == 2 and run["notQualified"] == 1 and run["undetermined"] == 1
        assert len(run["downloadFailures"]) == 1
        assert run["downloadFailures"][0]["excelRow"] == 6
        diagnostic = run["downloadFailures"][0]["reason"]
        assert "گزارش مراحل دریافت" in diagnostic and "HTTP 403" in diagnostic
        assert "SENSITIVE" not in diagnostic and "token=" not in diagnostic
        assert {r["jobinjaUrl"] for r in run["results"]} == {row[1] for row in rows[:4]}
        assert all(r["candidateName"] == "نام مشترک" and r["applicantId"] for r in run["results"])
        assert len({r["filename"].casefold() for r in run["results"]}) == 4
        report = client.get(f"/api/runs/{run_id}/report")
        workbook = load_workbook(BytesIO(report.content))
        metadata = list(workbook["دریافت رزومه‌ها"].values)
        assert len(metadata) == 6
        assert metadata[-1][5] == "دریافت نشد" and metadata[-1][8] is None
        assert metadata[-1][6] == diagnostic
        assert metadata[1][3] == rows[0][1]
        assert workbook["دریافت رزومه‌ها"]["D2"].hyperlink.target == rows[0][1]
        workbook.close()
        result = run["results"][0]
        review = client.post(f"/api/runs/{run_id}/review", json={"filename": result["filename"], "decision": "needs-review"})
        assert review.status_code == 200, review.text
        zipped = client.get(f"/api/runs/{run_id}/export/all")
        assert zipped.status_code == 200, zipped.text
        with ZipFile(BytesIO(zipped.content)) as archive:
            resumes = [n for n in archive.namelist() if "/resumes/" in n]
            assert len(resumes) == 4 and len(archive.namelist()) == 5 and archive.testzip() is None
            assert sorted(archive.read(n) for n in resumes) == sorted(originals)
    assert "SENSITIVE" not in caplog.text and "token=" not in caplog.text


def test_all_download_failures_have_report_and_no_evaluation(isolated, fake_http):
    responses, _ = fake_http
    responses[URL] = OSError("URL containing SENSITIVE failed")
    up = isolated.post("/api/jobinja/imports", files={"file": ("applicants.xlsx", excel([("Name", URL, "text")]))})
    import_id = up.json()["importId"]
    isolated.post(f"/api/jobinja/imports/{import_id}/run", json={"criteria": {"minYears": 5}})
    state = wait_for_import(isolated, import_id)
    run_id = state["runId"]
    run = isolated.get(f"/api/runs/{run_id}").json()
    assert run["results"] == [] and run["processed"] == 0 and len(run["downloadFailures"]) == 1
    assert isolated.get(f"/api/runs/{run_id}/report").status_code == 200
    assert isolated.get(f"/api/runs/{run_id}/export/all").status_code == 404


def test_partial_disk_write_is_not_evaluated_or_exported(isolated, monkeypatch):
    from backend.app import uploads
    monkeypatch.setattr(jobinja, "fetch_resume", lambda url: (docx("7+ years of professional experience"), ".docx"))
    def partial_write(directory, filename, data):
        (directory / "partial.docx").write_bytes(b"incomplete")
        raise OSError(errno.ENOSPC, "disk full")
    monkeypatch.setattr(uploads, "store_file", partial_write)
    up = isolated.post("/api/jobinja/imports", files={"file": ("applicants.xlsx", excel([("Name", URL, "text")]))})
    import_id = up.json()["importId"]
    isolated.post(f"/api/jobinja/imports/{import_id}/run", json={"criteria": {"minYears": 5}})
    state = wait_for_import(isolated, import_id)
    assert state["downloaded"] == 0 and state["failed"] == 1
    run = isolated.get(f"/api/runs/{state['runId']}").json()
    assert run["processed"] == 0 and not run["results"]
    assert "فضای دیسک" in run["downloadFailures"][0]["reason"]


def test_cancel_download_stops_remaining_applicants(isolated, monkeypatch):
    import threading
    entered, release = threading.Event(), threading.Event()
    seen = []
    def download(url):
        seen.append(url)
        entered.set()
        assert release.wait(3)
        return docx("7+ years of professional experience"), ".docx"
    monkeypatch.setattr(jobinja, "fetch_resume", download)
    rows = [("First", URL, "hyperlink"), ("Second", URL + "&applicant=2", "hyperlink")]
    up = isolated.post("/api/jobinja/imports", files={"file": ("applicants.xlsx", excel(rows))})
    import_id = up.json()["importId"]
    isolated.post(f"/api/jobinja/imports/{import_id}/run", json={"criteria": {"minYears": 5}})
    try:
        assert entered.wait(3)
        state = isolated.get(f"/api/jobinja/imports/{import_id}").json()
        assert state["state"] == "downloading" and state["completed"] == 0
        assert isolated.post(f"/api/jobinja/imports/{import_id}/cancel").status_code == 200
    finally:
        release.set()
    state = wait_for_import(isolated, import_id)
    assert seen == [URL] and state["downloaded"] == 1 and state["failed"] == 1
    run = isolated.get(f"/api/runs/{state['runId']}").json()
    assert run["processed"] == 1 and len(run["downloadFailures"]) == 1
    assert "لغو" in run["downloadFailures"][0]["reason"]


# --- Jobinja 403/404 that mean "the file or the share link is gone" ---------


def gone_page(title, text):
    """A Jobinja-style 'not found' page, exactly as the site serves it."""
    return f"<html><head><title>{title}</title></head><body><h1>{text}</h1></body></html>".encode()


def error_response(url, code, body=b"", headers=None):
    message = Message()
    message["Content-Type"] = "text/html; charset=UTF-8"
    for key, value in (headers or {}).items():
        message[key] = value
    return urllib.error.HTTPError(url, code, "error", message, BytesIO(body))


@pytest.mark.parametrize("stage, body, expected", [
    ("file", "<html><body>جابینجا این صفحه وجود ندارد یا آدرس آن عوض شده است!</body></html>".encode(), "موجود نیست"),
    ("file", "<html><body>404 page not found</body></html>".encode(), "موجود نیست"),
    ("page", "<html><body>جابینجا صفحه مورد نظر یافت نشد</body></html>".encode(), "وجود ندارد"),
])
def test_jobinja_gone_page_is_reported_with_a_precise_reason(fake_http, stage, body, expected):
    responses, requests = fake_http
    link = "https://jobinja.ir/user/NO-SYNTHETIC/download"
    if stage == "file":
        responses[URL] = (f'<a href="{link}" download>دانلود</a>'.encode(), "text/html")
        responses[link] = error_response(link, 403, body)
    else:
        responses[URL] = error_response(URL, 404, body)
    with pytest.raises(jobinja.JobinjaError) as error:
        jobinja.fetch_resume(URL)
    message = str(error.value)
    assert expected in message.splitlines()[0]
    assert "صفحهٔ «یافت نشد» جابینجا" in message
    assert "<html>" not in message and "SENSITIVE" not in message
    assert [r.full_url for r in requests] == ([URL, link] if stage == "file" else [URL])


def test_gzip_encoded_gone_page_is_still_classified(fake_http):
    responses, _ = fake_http
    link = "https://jobinja.ir/user/NO-SYNTHETIC/download"
    responses[URL] = (f'<a href="{link}" download>دانلود</a>'.encode(), "text/html")
    body = gzip.compress("<html><body>این صفحه وجود ندارد یا آدرس آن عوض شده است!</body></html>".encode())
    responses[link] = error_response(link, 403, body, {"Content-Encoding": "gzip"})
    with pytest.raises(jobinja.JobinjaError, match="موجود نیست"):
        jobinja.fetch_resume(URL)


def test_unclassified_403_keeps_the_plain_status_message(fake_http):
    responses, _ = fake_http
    link = "https://jobinja.ir/user/NO-SYNTHETIC/download"
    responses[URL] = (f'<a href="{link}" download>دانلود</a>'.encode(), "text/html")
    responses[link] = error_response(link, 403, b"")
    with pytest.raises(jobinja.JobinjaError, match="HTTP 403"):
        jobinja.fetch_resume(URL)


def test_gone_resume_stops_nothing_and_is_never_staged(isolated, fake_http):
    client = isolated
    responses, requests = fake_http
    rows = []
    for index, text in enumerate(["5+ years of professional experience", "7+ years of professional experience"]):
        url = URL + f"&candidate={index}"
        rows.append(("نام مشترک", url, "hyperlink"))
        download_url = "https://jobinja.ir/file?token=SENSITIVE&candidate=" + str(index)
        if index == 0:
            responses[url] = (f'<a href="{download_url.replace("&", "&amp;")}" download>دانلود</a>'.encode(), "text/html")
            responses[download_url] = error_response(
                download_url, 403,
                "<html><body>این صفحه وجود ندارد یا آدرس آن عوض شده است!</body></html>".encode())
        else:
            responses[url] = (f'<a href="{download_url.replace("&", "&amp;")}" download>دانلود</a>'.encode(), "text/html")
            responses[download_url] = (docx(text), "application/octet-stream")
    up = client.post("/api/jobinja/imports", files={"file": ("جابینجا.xlsx", excel(rows))})
    import_id = up.json()["importId"]
    client.post(f"/api/jobinja/imports/{import_id}/run", json={"criteria": {"minYears": 5}})
    state = wait_for_import(client, import_id)
    assert state["downloaded"] == 1 and state["completed"] == 2 and state["failed"] == 1
    run = client.get(f"/api/runs/{state['runId']}").json()
    assert run["state"] == "finished" and run["processed"] == 1
    assert "موجود نیست" in run["downloadFailures"][0]["reason"]
    # The missing one is a receipt error, never a screening verdict, and no HTML was stored.
    assert run["qualified"] + run["notQualified"] + run["undetermined"] == 1
    report = client.get(f"/api/runs/{state['runId']}/report")
    metadata = list(load_workbook(BytesIO(report.content))["دریافت رزومه\u200cها"].values)
    missing = [row for row in metadata[1:] if row[5] == "دریافت نشد"]
    assert len(missing) == 1 and missing[0][8] is None
    assert "موجود نیست" in missing[0][6]


def test_jobinja_viewer_and_zip_survive_staging_cleanup(isolated, fake_http):
    """Downloaded resumes stay viewable and exportable after staging is swept."""
    client = isolated
    responses, requests = fake_http
    rows, downloads = [], []
    # Two applicants share a name on purpose: written names must stay distinct.
    plan = [("متقاضی مشترک", True), ("متقاضی مشترک", True),
            ("نام کاملاً فارسی", True), ("نام دیگر", False)]
    for index, (name, should_download) in enumerate(plan):
        url = URL + f"&candidate={index}"
        link = "https://jobinja.ir/user/NO-SYNTHETIC/download?candidate=" + str(index)
        rows.append((name, url, "hyperlink"))
        responses[url] = (f'<a href="{link.replace("&", "&amp;")}" download>دانلود</a>'.encode(), "text/html")
        if should_download:
            data = pdf(f"{index}+ years of professional experience")
            downloads.append(data)
            responses[link] = (data, "application/octet-stream")
        else:
            responses[link] = error_response(link, 403,
                                             "<html><body>این صفحه وجود ندارد یا آدرس آن عوض شده است!</body></html>".encode())
    up = client.post("/api/jobinja/imports", files={"file": ("جابینجا.xlsx", excel(rows))})
    import_id = up.json()["importId"]
    client.post(f"/api/jobinja/imports/{import_id}/run", json={"criteria": {"minYears": 5}})
    state = wait_for_import(client, import_id)
    run_id = state["runId"]
    for _ in range(200):
        run = client.get(f"/api/runs/{run_id}").json()
        if run["state"] in ("finished", "failed", "cancelled"):
            break
        time.sleep(.02)
    assert run["state"] == "finished", run.get("error")
    assert run["processed"] == 3 and len(run["downloadFailures"]) == 1

    # The staged upload is temporary; the output folder is the durable copy.
    staging = Path(next(iter(runner._uploads.values())))
    uploads.sweep_stale(staging.parent, max_age_seconds=0)
    assert not staging.exists()

    written = sorted(Path(run["outputPath"]).iterdir())
    assert len({p.name for p in written}) == 4  # three resumes plus the report
    for result in run["results"]:
        preview = client.post(f"/api/runs/{run_id}/preview", json={"filename": result["filename"]})
        served = client.get(f"/api/runs/{run_id}/resumes/{result['filename']}")
        assert preview.status_code == 200, preview.text
        assert served.status_code == 200 and served.content, result["filename"]
    # A resume whose download failed is never served from a stand-in file.
    failed_name = "نام دیگر.pdf"
    missing = client.post(f"/api/runs/{run_id}/preview", json={"filename": failed_name})
    assert missing.status_code in (400, 404)
    assert client.get(f"/api/runs/{run_id}/resumes/{failed_name}").status_code == 404
    # Opening the viewer must not re-analyse or change the run.
    again = client.get(f"/api/runs/{run_id}").json()
    assert again["processed"] == 3 and again["results"] == run["results"]

    response = client.get(f"/api/runs/{run_id}/export/all")
    assert response.status_code == 200, response.text
    with ZipFile(BytesIO(response.content)) as archive:
        names = archive.namelist()
        resumes = [name for name in names if "/resumes/" in name]
        assert len(resumes) == 3 and "all-resumes/all-report.xlsx" in names
        assert len({name.casefold() for name in resumes}) == 3  # duplicate names kept apart
        assert archive.testzip() is None
        assert sorted(archive.read(name) for name in resumes) == sorted(downloads)
        workbook = load_workbook(BytesIO(archive.read("all-resumes/all-report.xlsx")))
        try:
            metadata = list(workbook["دریافت رزومه\u200cها"].values)
            assert len(metadata) == 5  # header plus one row per applicant
            assert [row[5] for row in metadata[1:]].count("دریافت نشد") == 1
        finally:
            workbook.close()


def test_jobinja_flow_works_under_a_persian_path_with_spaces(isolated, fake_http, monkeypatch, tmp_path):
    """Windows user folders are %LOCALAPPDATA% under the account name: Persian + spaces."""
    client = isolated
    responses, _ = fake_http
    root = tmp_path / "کاربر فارسی (test)"
    (root / "خروجی").mkdir(parents=True)
    monkeypatch.setattr(routes_runs.tempfile, "gettempdir", lambda: str(root))
    monkeypatch.setattr(runner, "user_data_dir", lambda: root / "خروجی")
    url = URL + "&candidate=persian"
    link = "https://jobinja.ir/user/NO-PERSIAN/download"
    data = pdf("6+ years of professional experience")
    responses[url] = (f'<a href="{link}" download>دانلود رزومه</a>'.encode(), "text/html")
    responses[link] = (data, "application/pdf")
    up = client.post("/api/jobinja/imports",
                     files={"file": ("جابینجا.xlsx", excel([("متقاضی فارسی", url, "hyperlink")]))})
    import_id = up.json()["importId"]
    client.post(f"/api/jobinja/imports/{import_id}/run", json={"criteria": {"minYears": 5}})
    state = wait_for_import(client, import_id)
    run_id = state["runId"]
    for _ in range(200):
        run = client.get(f"/api/runs/{run_id}").json()
        if run["state"] in ("finished", "failed", "cancelled"):
            break
        time.sleep(.02)
    assert run["state"] == "finished", run.get("error")
    assert "کاربر فارسی" in run["outputPath"] and (root / "خروجی").exists()
    served = client.get(f"/api/runs/{run_id}/resumes/{run['results'][0]['filename']}")
    assert served.status_code == 200 and served.content == data
    report = client.get(f"/api/runs/{run_id}/report")
    assert report.status_code == 200 and report.content[:2] == b"PK"
    archive = client.get(f"/api/runs/{run_id}/export/all")
    assert archive.status_code == 200
    with ZipFile(BytesIO(archive.content)) as z:
        assert sorted(z.namelist()) == ["all-resumes/all-report.xlsx",
                                        "all-resumes/resumes/" + run["results"][0]["filename"]]


@pytest.mark.parametrize("minimum", [None, -1, "nan", "inf", "abc"])
def test_only_finite_nonnegative_minimum_is_accepted(minimum):
    from fastapi import HTTPException
    with pytest.raises(HTTPException) as error:
        routes_runs._parse_criteria({"minYears": minimum})
    assert error.value.status_code == 400


def test_detected_skills_remain_informational_not_a_filter():
    from backend.app.models import ResumeFacts, ScreeningCriteria
    from backend.app.qualification import evaluate

    for skills in ([], ["React"], ["Vue"]):
        facts = ResumeFacts(filename="cv.pdf", declared_experience_years=7, skills_detected=skills)
        result = evaluate(facts, ScreeningCriteria(min_years=5))
        assert result.final_status.value == "Qualified" and result.skills_detected == skills
