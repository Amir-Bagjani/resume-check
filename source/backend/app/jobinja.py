"""Read applicant links from XLSX and stage real downloads for the existing runner."""
from __future__ import annotations

from dataclasses import dataclass, field
import gzip
from html.parser import HTMLParser
from http.cookiejar import CookieJar
from io import BytesIO
import ipaddress
import logging
from pathlib import Path
import re
import socket
import threading
import urllib.error
import urllib.parse
import urllib.request
import uuid
from zipfile import BadZipFile, ZipFile

from . import jobinja_challenge
from . import uploads
from .models import Applicant, ScreeningCriteria

logger = logging.getLogger(__name__)
MAX_EXCEL_BYTES = 10 * 1024 * 1024
MAX_HTML_BYTES = 2 * 1024 * 1024
MAX_APPLICANTS = 2000
_FORMULA = re.compile(r'^\s*=\s*HYPERLINK\s*\(\s*"((?:[^"\n]|"")*)"\s*[,;]', re.I)


class JobinjaError(Exception):
    """Safe Persian messages; never include the tokenized URL or response body."""


def _valid_url(url: str, domains: tuple[str, ...]) -> bool:
    try:
        parsed = urllib.parse.urlsplit(url)
        host = parsed.hostname or ""
        return (parsed.scheme == "https" and any(host == domain or host.endswith("." + domain) for domain in domains)
                and parsed.port in (None, 443) and parsed.username is None and parsed.password is None
                and not any(ord(c) <= 32 for c in url))
    except ValueError:
        return False


def valid_url(url: str) -> bool:
    return _valid_url(url, ("jobinja.ir",))


def valid_download_url(url: str) -> bool:
    return _valid_url(url, ("jobinja.ir", "jobinjacdn.com"))


def cell_url(cell) -> str:
    if cell.hyperlink and cell.hyperlink.target:
        return cell.hyperlink.target
    value = cell.value
    if not isinstance(value, str):
        return ""
    match = _FORMULA.match(value)
    if match:
        return match[1].replace('""', '"')
    return value.strip() if value.strip().lower().startswith(("https://", "http://")) else ""


def parse_excel(data: bytes) -> list[Applicant]:
    from openpyxl import load_workbook

    if len(data) > MAX_EXCEL_BYTES:
        raise JobinjaError("حجم Excel باید حداکثر ۱۰ مگابایت باشد.")
    try:
        with ZipFile(BytesIO(data)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 64 * 1024 * 1024:
                raise JobinjaError("حجم محتوای Excel بیش از حد مجاز است.")
        workbook = load_workbook(BytesIO(data), data_only=False, keep_links=False)
    except JobinjaError:
        raise
    except Exception:
        raise JobinjaError("فایل Excel معتبر نیست؛ فایل با پسوند XLSX انتخاب کنید.") from None
    try:
        sheet = workbook.worksheets[0]
        if sheet.max_row > MAX_APPLICANTS + 20 or sheet.max_column > 100:
            raise JobinjaError("Excel باید حداکثر ۲۰۰۰ متقاضی و ۱۰۰ ستون داشته باشد.")
        applicants = []
        name_column = None
        for row in sheet.iter_rows():
            if not any((cell.value is not None and str(cell.value).strip()) or cell.hyperlink for cell in row):
                continue
            links = list(dict.fromkeys(cell_url(cell) for cell in row if cell_url(cell)))
            if row[0].row == 1 and not links:
                labels = []
                for cell in row:
                    label = str(cell.value or "").replace("ي", "ی").replace("ك", "ک").replace("\u200c", " ").strip()
                    labels.append(label)
                    if label in ("نام", "نام متقاضی", "نام داوطلب", "نام کارجو", "نام و نام خانوادگی"):
                        name_column = cell.column - 1
                if name_column is not None or any(label in ("رزومه", "آدرس رزومه", "لینک رزومه", "مشاهده رزومه") for label in labels):
                    continue
            jobinja_links = [url for url in links if valid_url(url)]
            url = jobinja_links[0] if len(jobinja_links) == 1 else links[0] if len(links) == 1 else ""
            name = str(row[name_column].value or "") if name_column is not None else ""
            applicant = Applicant(uuid.uuid4().hex, row[0].row, name,
                                  jobinja_url=url)
            if not valid_url(applicant.jobinja_url):
                applicant.download_status = "failed"
                applicant.download_error = "این ردیف باید یک لینک معتبر رزومه جابینجا داشته باشد."
            applicants.append(applicant)
        if not applicants or len(applicants) > MAX_APPLICANTS:
            raise JobinjaError("Excel باید بین ۱ تا ۲۰۰۰ ردیف متقاضی داشته باشد.")
        return applicants
    finally:
        workbook.close()


class _DownloadLinks(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.links = []
        self.current = None

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == "a":
            self.current = [attrs.get("href", ""), "download" in attrs, attrs.get("aria-label", "") + " " + attrs.get("title", "")]

    def handle_data(self, data):
        if self.current is not None:
            self.current[2] += data

    def handle_endtag(self, tag):
        if tag == "a" and self.current is not None:
            href, download, label = self.current
            if href and (download or "دانلود" in label or "download" in label.lower()):
                self.links.append(href)
            self.current = None


def download_link(html: str, page_url: str, trace: list[str]) -> str:
    parser = _DownloadLinks()
    parser.feed(html)
    links = list(dict.fromkeys(urllib.parse.urljoin(page_url, link) for link in parser.links))
    trace.append(f"لینک‌های دانلود: {len(links)}؛ لینک‌های مجاز: {sum(valid_download_url(link) for link in links)}")
    scripts = len(re.findall(r"<script\b", html, re.I))
    trace.append(f"نشانه‌های HTML: script={scripts}؛ "
                 f"password={'password' in html.lower()}؛ captcha={'captcha' in html.lower()}؛ "
                 f"browser-challenge={_is_browser_challenge(html) or 'cf-chl' in html.lower() or 'challenge-platform' in html.lower()}")
    if len(links) != 1 or not valid_download_url(links[0]):
        raise JobinjaError("لینک دانلود رزومه در صفحه دریافت‌شده پیدا نشد؛ لینک را در مرورگر همین دستگاه باز کنید و ورود به حساب، اعتبار لینک یا صفحهٔ تأیید مرورگر را بررسی کنید.")
    return links[0]


class _NoRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        return None


def _check_destination(url: str) -> int:
    if not valid_download_url(url):
        raise JobinjaError("آدرس دریافت رزومه معتبر نیست.")
    addresses = socket.getaddrinfo(urllib.parse.urlsplit(url).hostname, 443, type=socket.SOCK_STREAM)
    if not addresses or any(not ipaddress.ip_address(item[4][0]).is_global for item in addresses):
        raise JobinjaError("آدرس دریافت رزومه قابل دسترسی نیست.")
    return len(addresses)


def _is_browser_challenge(html: str) -> bool:
    return "__arcsjs" in html and "location.reload" in html


# Jobinja serves its own "page not found" page (403/404) when a shared link or a
# stored CV is gone. The markers are compared against the response body, which is
# never logged or returned.
_GONE_MARKERS = ("وجود ندارد", "یافت نشد", "آدرس آن عوض شده", "page not found")
_GONE_MESSAGES = {
    "page": "صفحهٔ اشتراک این رزومه روی جابینجا دیگر وجود ندارد؛ لینک منقضی، حذف یا اشتباه است.",
    "file": "فایل رزومه روی جابینجا موجود نیست؛ رزومه حذف شده یا دسترسی عمومی به آن بسته شده است.",
}


def _classify_error(code: int, body: bytes, headers=None, stage: str = "page") -> tuple[str, bool]:
    """Describe an error response; returns (trace note, resource is gone)."""
    if code not in (401, 403, 404):
        return "", False
    encoding = (headers.get("Content-Encoding", "") if headers else "").lower()
    if "gzip" in encoding:
        try:
            body = gzip.GzipFile(fileobj=BytesIO(body)).read(8192)
        except OSError:
            body = b""
    text = body[:8192].decode("utf-8", errors="replace")
    gone = any(marker in text for marker in _GONE_MARKERS)
    note = (f"پاسخ {code}: " + ("صفحهٔ «یافت نشد» جابینجا" if gone else "بدون پیام قابل تشخیص")
            + "؛ مسیر و Query و متن پاسخ مخفی")
    return note, gone


def _fetch(url: str, opener, trace: list[str], referer: str = "", stage: str = "page") -> tuple[bytes, str, str]:
    # Check every redirect before opening it; only Jobinja and its CDN are allowed.
    for _ in range(6):
        if not valid_download_url(url):
            trace.append("مقصد درخواست خارج از دامنه‌های مجاز یا آدرس نامعتبر است؛ درخواست ارسال نشد.")
            raise JobinjaError("آدرس دریافت رزومه معتبر نیست.")
        trace.append(f"GET https://{urllib.parse.urlsplit(url).hostname}/ [مسیر و Query مخفی]")
        try:
            address_count = _check_destination(url)
            trace.append(f"DNS: {address_count} آدرس عمومی؛ timeout=20s")
            headers = {
                "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/131.0.0.0 Safari/537.36",
                "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
                "Accept-Language": "fa-IR,fa;q=0.9,en-US;q=0.8,en;q=0.7",
                "Accept-Encoding": "gzip",
                "Upgrade-Insecure-Requests": "1",
            }
            if referer:
                source = urllib.parse.urlsplit(referer)
                target = urllib.parse.urlsplit(url)
                # Browser default: signed query strings stay on the same origin.
                headers["Referer"] = (referer if source.netloc == target.netloc
                                      else f"{source.scheme}://{source.netloc}/")
            request = urllib.request.Request(url, headers=headers)
            with opener.open(request, timeout=20) as response:
                content_type = response.headers.get_content_type()
                encoding = response.headers.get("Content-Encoding", "identity").lower()
                trace.append(f"HTTP {response.status}؛ Content-Type={content_type[:100]}؛ Content-Encoding={encoding[:30]}")
                limit = MAX_HTML_BYTES if content_type == "text/html" else uploads.MAX_FILE_BYTES
                data = response.read(limit + 1)
                trace.append(f"حجم پاسخ: {len(data)} bytes؛ سقف: {limit} bytes")
                if len(data) > limit:
                    raise JobinjaError("حجم فایل دریافت‌شده بیش از حد مجاز است.")
                if encoding == "gzip":
                    with gzip.GzipFile(fileobj=BytesIO(data)) as compressed:
                        data = compressed.read(limit + 1)
                    trace.append(f"حجم پس از gzip: {len(data)} bytes")
                    if len(data) > limit:
                        raise JobinjaError("حجم فایل دریافت‌شده بیش از حد مجاز است.")
                return data, content_type, url
        except urllib.error.HTTPError as exc:
            location = exc.headers.get("Location")
            code = exc.code
            try:
                body = exc.read(8192)
            except Exception:
                body = b""
            trace.append(f"HTTP {code}؛ redirect={'yes' if location else 'no'}")
            exc.close()
            if code in (301, 302, 303, 307, 308) and location:
                url = urllib.parse.urljoin(url, location)
                continue
            note, gone = _classify_error(code, body, exc.headers, stage)
            if note:
                trace.append(note)
            if gone:
                # The site answered with its own "not found" page: the link or the
                # file is really unavailable, so it is reported instead of retried.
                raise JobinjaError(_GONE_MESSAGES[stage]) from None
            raise JobinjaError(f"دریافت رزومه از جابینجا ناموفق بود (HTTP {code})؛ دسترسی لینک را بررسی کنید.") from None
        except JobinjaError:
            raise
        except Exception as exc:
            reason = exc.reason if isinstance(exc, urllib.error.URLError) else exc
            trace.append(f"خطای شبکه/خواندن: {type(reason).__name__}؛ errno={getattr(reason, 'errno', None)}؛ "
                         f"TLS verify_code={getattr(reason, 'verify_code', None)}")
            raise JobinjaError("ارتباط با جابینجا برقرار نشد؛ اتصال اینترنت را بررسی کنید.") from None
    raise JobinjaError("تعداد تغییر مسیرهای دانلود بیش از حد مجاز است.")


def fetch_resume(url: str) -> tuple[bytes, str]:
    trace = ["درخواست: Chrome User-Agent؛ Accept=HTML؛ زبان=fa-IR؛ encoding=gzip؛ نشست کوکی مستقل"]
    try:
        return _fetch_resume(url, trace)
    except JobinjaError as exc:
        raise JobinjaError(str(exc) + "\n\nگزارش مراحل دریافت:\n" + "\n".join(trace)) from None


def _fetch_resume(url: str, trace: list[str]) -> tuple[bytes, str]:
    cookies = CookieJar()
    opener = urllib.request.build_opener(_NoRedirect(), urllib.request.HTTPCookieProcessor(cookies))
    trace.append(f"Proxy سیستم: {'فعال' if any(isinstance(h, urllib.request.ProxyHandler) and h.proxies for h in getattr(opener, 'handlers', [])) else 'غیرفعال'}")
    data, media, page_url = _fetch(url, opener, trace)
    if media == "text/html" or data.lstrip().lower().startswith((b"<!doctype html", b"<html")):
        html = data.decode("utf-8", errors="replace")
        if _is_browser_challenge(html):
            trace.append("صفحهٔ انتقال JavaScript شناسایی شد (__arcsjs)")
            try:
                # ponytail: the challenge is pure arithmetic, so no browser is needed.
                jobinja_challenge.clear(html, cookies, trace)
                data, _, page_url = _fetch(page_url, opener, trace)
                html = data.decode("utf-8", errors="replace")
                if _is_browser_challenge(html):
                    raise jobinja_challenge.ChallengeError("challenge page returned again")
            except jobinja_challenge.ChallengeError as exc:
                trace.append(f"نیاز به اجرای مرورگر واقعی: {exc}")
                from .jobinja_browser import render_resume_page
                html, page_url = render_resume_page(page_url, cookies, trace)
        link = download_link(html, page_url, trace)
        trace.append("لینک دانلود استخراج شد؛ دریافت فایل با همان نشست کوکی و Referer")
        data, _, _ = _fetch(link, opener, trace, referer=page_url, stage="file")
    trace.append("اعتبارسنجی محتوای PDF/DOCX")
    if b"%PDF-" in data[:1024]:
        from pypdf import PdfReader
        try:
            pdf = PdfReader(BytesIO(data))
            if not pdf.is_encrypted and not pdf.pages:
                raise ValueError("Unreadable PDF")
        except Exception:
            raise JobinjaError("فایل PDF دریافت‌شده معتبر یا قابل خواندن نیست.") from None
        return data, ".pdf"
    try:
        with ZipFile(BytesIO(data)) as archive:
            if ("[Content_Types].xml" not in archive.namelist() or "word/document.xml" not in archive.namelist()
                    or sum(info.file_size for info in archive.infolist()) > 64 * 1024 * 1024
                    or archive.testzip() is not None):
                raise BadZipFile()
        from docx import Document
        Document(BytesIO(data))
    except Exception:
        raise JobinjaError("فایل دریافت‌شده یک PDF یا DOCX معتبر نیست.") from None
    return data, ".docx"


@dataclass
class ImportJob:
    import_id: str
    folder: Path
    label: str
    applicants: list[Applicant]
    state: str = "ready"
    completed: int = 0
    downloaded: int = 0
    run_id: str | None = None
    cancelled: bool = False
    error: str = ""
    lock: threading.Lock = field(default_factory=threading.Lock, repr=False)


_imports: dict[str, ImportJob] = {}


def import_state(job: ImportJob) -> dict:
    with job.lock:
        return {"importId": job.import_id, "folderName": job.label, "state": job.state,
                "total": len(job.applicants), "completed": job.completed, "downloaded": job.downloaded,
                "failed": sum(a.download_status == "failed" for a in job.applicants),
                "runId": job.run_id, "error": job.error}


def _download(job: ImportJob, criteria: ScreeningCriteria) -> None:
    from . import runner

    total_bytes = 0
    try:
        job.folder.mkdir(parents=True, exist_ok=True)
        for applicant in job.applicants:
            with job.lock:
                cancelled = job.cancelled
            if applicant.download_status == "failed":
                pass
            elif cancelled:
                applicant.download_status = "failed"
                applicant.download_error = "دریافت این رزومه لغو شد."
            else:
                try:
                    data, extension = fetch_resume(applicant.jobinja_url)
                    if total_bytes + len(data) > uploads.MAX_TOTAL_BYTES:
                        raise JobinjaError("حجم کل رزومه‌ها بیش از حد مجاز است.")
                    filename = uploads.safe_filename(applicant.name or f"متقاضی-{applicant.excel_row}") + extension
                    stored = uploads.store_file(job.folder, filename, data)
                    applicant.filename = stored.as_posix()
                    applicant.download_status = "downloaded"
                    total_bytes += len(data)
                except Exception as exc:
                    applicant.download_status = "failed"
                    applicant.download_error = str(exc) if isinstance(exc, JobinjaError) else "ذخیره فایل رزومه ناموفق بود؛ فضای دیسک و دسترسی را بررسی کنید."
                    logger.warning("Applicant download failed: import=%s row=%d applicant=%s", job.import_id, applicant.excel_row, applicant.applicant_id)
            with job.lock:
                job.completed += 1
                job.downloaded += applicant.download_status == "downloaded"
        runner.register_upload(job.folder)
        run = runner.create_run(str(job.folder), criteria, applicants=job.applicants)
        with job.lock:
            job.run_id = run.run_id
            job.state = "screening"
    except Exception:
        # Exception messages can contain URLs: log only stable identifiers here.
        logger.error("Jobinja import could not finish: import=%s", job.import_id)
        with job.lock:
            job.state = "failed"
            job.error = "آماده‌سازی گزارش و بررسی رزومه‌ها ناموفق بود؛ فضای دیسک و دسترسی را بررسی کنید."


def start_import(job: ImportJob, criteria: ScreeningCriteria, output_folder: str | None = None) -> None:
    from . import runner

    with job.lock:
        if job.state != "ready":
            raise JobinjaError("دریافت این فایل قبلاً شروع شده است.")
        runner.set_output_target(job.folder.name, output_folder)
        job.state = "downloading"
    thread = threading.Thread(target=_download, args=(job, criteria), daemon=True)
    thread.start()
